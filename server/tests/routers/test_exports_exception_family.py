"""exports 异常族端点面镜像测试：exp-hygiene-20260930 回炉轮1（R1 映射 3/3）+est-20261001（R2 翻面）。

输入:  POST /api/exports/audit·calcbook（真实链路：routers→services→jobs/worker→core）
输出:  异常族契约断言（flows 三件逐族 422 映射/estimate 真产物+边车点亮）
"""

# ══════════════════════════════════════════════════════════════════
# 规格说明（exp-hygiene-20260930 回炉轮1——门一 k2-W5+d1-W1/W2/R2
#   k2-W2+d1-N7；test_exports_audit.py 484 行顶 §2 预算墙拆件，公共
#   前置束经命名空间包导入共享〔pythonpath=. 先例〕）
#
# 【覆盖面】
#   - R1：H2 main_lib._EXCEPTION_STATUS 映射面 **3/3 参数化**——真实链路
#     （TestClient POST /api/exports/audit 单产物）+monkeypatch flows.
#     audit_render_flow 逐族注入→422+error_type 类名保真+异常即零新
#     产物入册（替换 test_exports_audit.py 原 dict 轮询版）；
#   - R2（est-20261001 D5 翻面）：estimate 真产物面——混装批 {calcbook,
#     audit,estimate}：estimate 项真出 .xlsx（真渲染链 flows.
#     estimate_render_flow）+边车点亮（H1 面复验）+零 failures（原
#     「501 未就绪项级失败/零产物零边车」诚实面断言随 501 收口反转；
#     批零替身原则沿承）。
# 【替身口径】仅 monkeypatch flows.audit_render_flow（R1 注入面）；
#   R2 零替身（estimate 成功=真渲染链）。
# ══════════════════════════════════════════════════════════════════

from __future__ import annotations

import importlib
import os

import pytest
from fastapi import status
from openpyxl import load_workbook

from tests.routers.test_exports_audit import _project_with_result, _wait_task_terminal

_flows = importlib.import_module("waterprint.flows")

pytestmark = [pytest.mark.anyio]

# R1：flows 异常族三名（参数化成员=main_lib._EXCEPTION_STATUS 三条映射逐条对拍）
_FAMILY: tuple[str, ...] = ("InvalidFlowError", "InvalidAuditError", "InvalidAuditPathError")


@pytest.mark.anyio
@pytest.mark.parametrize("exc_name", _FAMILY)
async def test_audit_render_exception_family_maps_422_wiring(  # type: ignore[no-untyped-def]
    client, monkeypatch, exc_name: str
) -> None:
    """R1（回炉轮1 k2-W5+d1-W1/W2）：H2 映射面 3/3——真实链路单产物
    POST /api/exports/audit+flows.audit_render_flow 逐族注入→422
    Unprocessable Content+error_type 类名保真（GR-11 参数族——用户
    输入域非服务端 500；未映射成员=裸 500 炸穿面逐条钉死）。"""
    exc_type = getattr(_flows, exc_name)

    def _raise(project, plant, out):  # type: ignore[no-untyped-def]
        raise exc_type(f"injected {exc_name} (R1 family mapping)")

    monkeypatch.setattr(_flows, "audit_render_flow", _raise)
    project_id, _task_id = await _project_with_result(client)
    before = len(
        (await client.get("/api/exports", params={"project_id": project_id})).json()
    )
    response = await client.post("/api/exports/audit", json={"project_id": project_id})
    assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT  # 映射在册
    assert response.json()["error_type"] == exc_name  # 类名保真（诊断面）
    assert "injected" in str(response.json()["detail"])  # 异常消息透传（R1 注入锚）
    after = len(
        (await client.get("/api/exports", params={"project_id": project_id})).json()
    )
    assert after == before  # 异常即零新产物入册


@pytest.mark.anyio
async def test_mixed_batch_estimate_true_product_and_sidecar_wiring(  # type: ignore[no-untyped-def]
    client, test_settings
) -> None:
    """R2 翻面（est-20261001 D5）：estimate 真产物面——混装批 {calcbook,
    audit,estimate}：estimate 项真出 .xlsx（真渲染链）+边车点亮（H1 面
    复验）+零 failures（原 501 未就绪期「项级失败/零产物零边车」诚实面
    断言随 501 收口反转；批零替身原则沿承）。"""
    project_id, _task_id = await _project_with_result(client)
    batch = await client.post(
        "/api/exports/calcbook",
        json={
            "project_id": project_id,
            "options": {"items": [
                {"kind": "calcbook", "condition_key": "design"},
                {"kind": "audit"},
                {"kind": "estimate", "condition_key": "design"},
            ]},
        },
    )
    assert batch.status_code == status.HTTP_200_OK  # 批量转任务句柄 JSON
    done = await _wait_task_terminal(client, str(batch.json()["task_id"]))
    assert done["state"] == "done"
    assert len(done["result"]["files"]) == 3  # 三 kind 全成功（estimate 真产物）
    assert list(done["result"]["failures"]) == []  # 零 failures（含 estimate 项）
    listing = sorted(os.listdir(test_settings.exports_dir))
    estimate_products = [n for n in listing if "-estimate-" in n and n.endswith(".xlsx")]
    assert len(estimate_products) == 1  # estimate 产物各一（命名段锚）
    assert "-estimate-design-" in estimate_products[0]  # 工况段=归一 design
    assert len([n for n in listing if n.endswith(".meta.json")]) == 3  # 三边车各一
    assert (test_settings.exports_dir / f"{estimate_products[0]}.meta.json").is_file()
    workbook = load_workbook(test_settings.exports_dir / estimate_products[0])
    assert workbook.sheetnames == ["概算汇总表", "指标校核表"]  # 真渲染链锚点
    metas = await client.get("/api/exports", params={"project_id": project_id})
    kinds = {row["kind"] for row in metas.json()}
    assert {"calcbook", "audit", "estimate"} <= kinds  # 三 kind 均入注册表
    estimate_row = next(row for row in metas.json() if row["kind"] == "estimate")
    downloaded = await client.get(f"/api/exports/{estimate_row['file_name']}")
    assert downloaded.status_code == status.HTTP_200_OK  # 边车在场=可下载
    assert downloaded.content == (  # 字节==落盘产物（非重渲染）
        test_settings.exports_dir / estimate_row["file_name"]
    ).read_bytes()
