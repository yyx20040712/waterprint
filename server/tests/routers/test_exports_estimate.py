"""exports estimate 端点镜像测试：est-20261001 收口批（501→200 xlsx 真产物）。

输入:  waterprint_server.routers.exports /estimate 端点（service→flows 全链）
输出:  estimate 导出行为契约断言（200 xlsx/幂等命名+字节/工况归一与 422
       闸/未知工况 422/混装批零路由键继承）
"""

# ══════════════════════════════════════════════════════════════════
# 规格说明（est-20261001 estimate 501 收口批·server 通道笔——
#   test_exports_audit.py 同族形态）
#
# 【覆盖面】
#   - 单产物正向：POST /api/exports/estimate → 200 xlsx 媒体型+锚点内容
#     （概算汇总表/指标校核表/抬头三元组串）+.xlsx 后缀+落盘真件+幂等
#     命名（同输入同名）+幂等重导出 sha256 恒等（§七探针断言位）+
#     GET 下载 200 字节==落盘产物；
#   - 工况归一（D4）：bare POST 与显式 condition_key="design" 同名同字节
#     =幂等（注册表单行）；condition_key=estimate 合法语义项（audit 反向）；
#   - 未知工况：condition_key 不在结果集 → 422 InvalidFlowError（消息含
#     可用工况集——渲染流包装，禁静默取首档）+零落盘；
#   - 选项闸（D4 audit 闸镜像）：item 级/纯批批级 unit_id、六路由键任一
#     携带（判据类型收死——unit_id=123 同拒）→ 422 且零落盘；
#   - 混装批零继承：dxf+estimate items（批级 unit_id）→ dxf 项继承 unit
#     分量、estimate 项全路由键置空（命名无 unit 段）+任务终态 done
#     双产物（worker 分流实证）。
#   - 回炉轮1（rework-est-20261001-r1）：R1 非 estimate 项显式 null 行为锚
#     （dxf condition_key:null 命名段维持既有形态+注册表 condition 空串
#     ——零物化零行为变）+estimate 显式空串归一 design 同名；R2 端点工况
#     缺省序（item 自有→端点值→design——items-present 不静默吞端点值）；
#     R4 双处缺省常量等价锁（jobs/services 层序隔离——漂移即红）；R6
#     批量句柄 condition_key 回显四面同源（=items[0] 归一值）。
# 【替身口径】零替身——estimate 渲染走真 flows.estimate_render_flow
#   （本批被测接线面即分流正门；audit 批同款）。
# ══════════════════════════════════════════════════════════════════

from __future__ import annotations

import importlib
import os
from hashlib import sha256

import pytest
from fastapi import status
from openpyxl import load_workbook

from tests.routers.test_exports_audit import _project_with_result, _wait_task_terminal

_mod = importlib.import_module("waterprint_server.routers.exports")
router = getattr(_mod, "router")

pytestmark = [
    pytest.mark.skipif(
        router is None,
        reason="实现未就绪：waterprint_server.routers.exports（服务层 M2/M3）",
    ),
    pytest.mark.anyio,
]


def _disposition_name(response) -> str:  # type: ignore[no-untyped-def]
    """Content-Disposition 文件名解析（test_exports_audit.py 同款）。"""
    disposition = str(response.headers.get("content-disposition", ""))
    return disposition.rsplit('filename="', maxsplit=1)[-1].rstrip('"')


@pytest.mark.anyio
async def test_estimate_single_product_xlsx_flow_wiring(  # type: ignore[no-untyped-def]
    client, test_settings
) -> None:
    """单产物正向：200 xlsx+锚点（两表名+抬头三元组）+.xlsx 后缀+落盘真件
    +幂等命名+幂等重导出 sha256 恒等+下载闸 200 字节==落盘件。"""
    project_id, task_id = await _project_with_result(client)
    engine_version = (await client.get(f"/api/calc/tasks/{task_id}")).json()[
        "result"
    ]["engine_version"]
    fresh = await client.post("/api/exports/estimate", json={"project_id": project_id})
    assert fresh.status_code == status.HTTP_200_OK
    assert fresh.headers["content-type"].startswith("application/vnd")  # xlsx 媒体型
    assert _disposition_name(fresh).endswith(".xlsx")  # kind 后缀映射
    file_name = _disposition_name(fresh)
    product = test_settings.exports_dir / file_name
    assert product.is_file()  # 落盘真件
    workbook = load_workbook(product)
    assert workbook.sheetnames == ["概算汇总表", "指标校核表"]  # 渲染契约锚点
    text = "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )
    assert str(engine_version) in text  # 三元组版本串在场（头部自证）
    assert "工程概算总计" in text
    downloaded = await client.get(f"/api/exports/{file_name}")
    assert downloaded.status_code == status.HTTP_200_OK  # 边车在场解锁下载面
    assert downloaded.content == product.read_bytes()  # 字节==落盘产物
    metas = await client.get("/api/exports", params={"project_id": project_id})
    rows = [meta for meta in metas.json() if meta["kind"] == "estimate"]
    assert len(rows) == 1 and rows[0]["file_name"] == file_name
    assert rows[0]["condition_key"] == "design"  # 缺省归一回显（D4 四面同源）
    again = await client.post("/api/exports/estimate", json={"project_id": project_id})
    assert again.status_code == status.HTTP_200_OK
    assert _disposition_name(again) == file_name  # 幂等命名（同输入同名）
    assert sha256(product.read_bytes()).hexdigest() == sha256(
        (test_settings.exports_dir / _disposition_name(again)).read_bytes()
    ).hexdigest()  # 幂等重导出 sha256 恒等（§七探针——确定性渲染）


@pytest.mark.anyio
async def test_estimate_condition_default_design_idempotent_wiring(client) -> None:  # type: ignore[no-untyped-def]
    """工况归一（D4）：bare POST 与显式 condition_key="design" 同名同字节
    =幂等（注册表单行）；condition_key=estimate 合法面（audit 语义反向）。"""
    project_id, _task_id = await _project_with_result(client)
    bare = await client.post("/api/exports/estimate", json={"project_id": project_id})
    assert bare.status_code == status.HTTP_200_OK
    explicit = await client.post(
        "/api/exports/estimate",
        json={"project_id": project_id, "condition_key": "design"},
    )
    assert explicit.status_code == status.HTTP_200_OK  # condition_key 合法（不 422）
    empty_string = await client.post(  # 回炉 R1②：items 显式空串归一 design 同名
        "/api/exports/estimate",
        json={"project_id": project_id,
              "options": {"items": [{"kind": "estimate", "condition_key": ""}]}},
    )
    assert empty_string.status_code == status.HTTP_200_OK
    assert _disposition_name(bare) == _disposition_name(explicit)  # 同名（归一单点）
    assert _disposition_name(bare) == _disposition_name(empty_string)  # 空串==design
    metas = await client.get("/api/exports", params={"project_id": project_id})
    rows = [meta for meta in metas.json() if meta["kind"] == "estimate"]
    assert len(rows) == 1  # 幂等覆盖（注册表不重复登记）


@pytest.mark.anyio
async def test_estimate_unknown_condition_422_wiring(  # type: ignore[no-untyped-def]
    client, test_settings
) -> None:
    """未知工况：不在结果集 → 422 InvalidFlowError（消息含可用工况集——
    渲染流包装 cost 域异常，禁静默取首档）+零落盘。"""
    project_id, _task_id = await _project_with_result(client)
    exports_dir = test_settings.exports_dir
    before = sorted(os.listdir(exports_dir))
    response = await client.post(
        "/api/exports/estimate",
        json={"project_id": project_id, "condition_key": "no_such_condition"},
    )
    assert response.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    body = response.json()
    assert body["error_type"] == "InvalidFlowError"  # H2 flows 族映射单一真源
    assert "可用工况" in str(body["detail"])  # 指路消息（禁静默取首档）
    assert "design" in str(body["detail"])
    assert sorted(os.listdir(exports_dir)) == before  # 拒绝即零落盘


@pytest.mark.anyio
async def test_estimate_unit_id_and_route_options_rejected_422_wiring(  # type: ignore[no-untyped-def]
    client, test_settings
) -> None:
    """选项闸（D4 audit 闸镜像）：estimate 项 unit_id/六路由键任一携带
    （item 级/纯批批级/非字符串形态）→ 422 且零落盘（概算=全厂整厂产物
    不分单元——工况选择走 condition_key）。"""
    project_id, _task_id = await _project_with_result(client)
    exports_dir = test_settings.exports_dir
    before = sorted(os.listdir(exports_dir))
    item_unit = await client.post(
        "/api/exports/estimate",
        json={"project_id": project_id,
              "options": {"items": [{"kind": "estimate", "unit_id": "municipal_cass"}]}},
    )
    assert item_unit.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    assert item_unit.json()["error_type"] == "InvalidExportRequestError"
    assert "unit_id" in str(item_unit.json()["detail"])  # 中文指路文案
    batch_unit = await client.post(  # 纯 estimate 批批级 unit_id → 422
        "/api/exports/estimate",
        json={"project_id": project_id,
              "options": {"unit_id": "municipal_cass", "items": [{"kind": "estimate"}]}},
    )
    assert batch_unit.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    item_sheet = await client.post(  # item 级 sheet（六路由键族）→ 422
        "/api/exports/estimate",
        json={"project_id": project_id,
              "options": {"items": [{"kind": "estimate", "sheet": "profile"}]}},
    )
    assert item_sheet.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    assert "sheet" in str(item_sheet.json()["detail"])
    typed_unit = await client.post(  # unit_id=123（非字符串）→ 422（判据收死）
        "/api/exports/estimate",
        json={"project_id": project_id,
              "options": {"items": [{"kind": "estimate", "unit_id": 123}]}},
    )
    assert typed_unit.status_code == status.HTTP_422_UNPROCESSABLE_CONTENT
    assert sorted(os.listdir(exports_dir)) == before  # 拒绝即零落盘


@pytest.mark.anyio
async def test_estimate_mixed_batch_no_inherit_and_worker_branch_wiring(  # type: ignore[no-untyped-def]
    client, test_settings
) -> None:
    """混装批零继承+批量路径：dxf+estimate items（批级 unit_id）→dxf 项继承
    unit 分量、estimate 项全路由键置空（命名无 unit 段+.xlsx）；任务终态
    done 双产物（dxf 魔面+xlsx 真渲染——worker 分流 flows 实证）。"""
    project_id, _task_id = await _project_with_result(client)
    batch = await client.post(
        "/api/exports/dxf",
        json={
            "project_id": project_id,
            "options": {
                "unit_id": "municipal_cass",  # 批级——dxf 项继承/estimate 项置空
                "items": [
                    {"kind": "dxf", "condition_key": "design"},
                    {"kind": "estimate", "condition_key": "design"},
                ],
            },
        },
    )
    assert batch.status_code == status.HTTP_200_OK  # 批量转任务句柄 JSON
    assert str(batch.json()["condition_key"]) == "design"  # R6：句柄回显=items[0] 归一值
    done = await _wait_task_terminal(client, str(batch.json()["task_id"]))
    assert done["state"] == "done" and len(done["result"]["files"]) == 2
    assert list(done["result"]["failures"]) == []  # 零 failures（含 estimate 项）
    listing = sorted(os.listdir(test_settings.exports_dir))
    xlsx_names = [name for name in listing if name.endswith(".xlsx")]
    dxf_names = [name for name in listing if name.endswith(".dxf")]
    assert len(dxf_names) == 1 and "-dxf-municipal_cass-design-" in dxf_names[0]
    assert len(xlsx_names) == 1  # estimate 产物各一（模板夹具不入 exports_dir）
    assert "-estimate-design-" in xlsx_names[0]  # 工况段=归一 design
    assert "municipal_cass" not in xlsx_names[0]  # estimate 项零路由键继承
    workbook = load_workbook(test_settings.exports_dir / xlsx_names[0])
    assert workbook.sheetnames == ["概算汇总表", "指标校核表"]  # 真 flows 渲染
    assert b"AC1032" in (test_settings.exports_dir / dxf_names[0]).read_bytes()[:512]
    metas = await client.get("/api/exports", params={"project_id": project_id})
    estimate_row = next(  # R6：注册表 condition=归一值（四面同源）
        row for row in metas.json() if row["kind"] == "estimate"
    )
    assert estimate_row["condition_key"] == "design"


@pytest.mark.anyio
async def test_non_estimate_null_condition_prebatch_behavior_anchor_wiring(  # type: ignore[no-untyped-def]
    client, test_settings
) -> None:
    """R1①（回炉轮1——双席共指）：非 estimate 项显式 null 零物化零行为变
    ——dxf item condition_key:null → 批前行为锁死：渲染静默首档（200）
    +命名段维持既有形态（null→str 物化段"None"为命名读取处历史形态）
    +注册表 condition 空串（meta 读取处 or-归一）。曾反转面：归一层
    str(null)="None" 串致渲染入参/边车携带非空工况（回炉修复锚）。"""
    project_id, _task_id = await _project_with_result(client)
    response = await client.post(
        "/api/exports/dxf",
        json={"project_id": project_id,
              "options": {"items": [{"kind": "dxf", "condition_key": None}]}},
    )
    assert response.status_code == status.HTTP_200_OK  # 批前行为：不 422/500
    file_name = _disposition_name(response)
    assert file_name.endswith(".dxf")
    assert "-None-" in file_name  # 命名段既有形态（命名读取处 str(null) 历史行为）
    metas = await client.get("/api/exports", params={"project_id": project_id})
    dxf_rows = [meta for meta in metas.json() if meta["kind"] == "dxf"]
    assert len(dxf_rows) == 1 and dxf_rows[0]["file_name"] == file_name
    assert dxf_rows[0]["condition_key"] == ""  # 注册表空串（meta or-归一）


@pytest.mark.anyio
async def test_estimate_endpoint_condition_fallback_chain_wiring(  # type: ignore[no-untyped-def]
    client, test_settings
) -> None:
    """R2（回炉轮1 d1-F2）：estimate 项缺省序=item 自有→端点值→"design"
    （端点参数本质=批级意图——SVRB D1 options.unit_id 同族；items-present
    不静默吞端点工况）：端点 condition_key="avg"+项无自有 → 渲染 avg+命名
    含 avg+注册表回显 avg。"""
    project_id, _task_id = await _project_with_result(client)
    response = await client.post(
        "/api/exports/estimate",
        json={
            "project_id": project_id,
            "condition_key": "avg",
            "options": {"items": [{"kind": "estimate"}]},
        },
    )
    assert response.status_code == status.HTTP_200_OK
    file_name = _disposition_name(response)
    assert "-estimate-avg-" in file_name  # 端点值入命名（不被 items-present 吞）
    assert (test_settings.exports_dir / file_name).is_file()
    metas = await client.get("/api/exports", params={"project_id": project_id})
    rows = [meta for meta in metas.json() if meta["kind"] == "estimate"]
    assert len(rows) == 1 and rows[0]["condition_key"] == "avg"  # 回显四面同源


def test_estimate_default_condition_constants_equivalence_lock() -> None:
    """R4（回炉轮1 k2-N2）：双处缺省常量等价锁——jobs 不可 import services
    （层序），字面量双处声明的漂移由本断言即时红。"""
    from waterprint_server.jobs.export_render import (
        _ESTIMATE_DEFAULT_CONDITION as _JOBS_DEFAULT,
    )
    from waterprint_server.services.exports import (
        _ESTIMATE_DEFAULT_CONDITION as _SERVICES_DEFAULT,
    )

    assert _JOBS_DEFAULT == _SERVICES_DEFAULT == "design"
