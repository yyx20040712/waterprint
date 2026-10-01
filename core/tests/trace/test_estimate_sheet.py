"""estimate xlsx 渲染件镜像测试：est-20261001（概算 501 收口批——案乙直写）。

输入:  waterprint.trace.estimate_sheet.render_estimate_xlsx +
       waterprint.flows.estimate_render_flow（golden 计算链真值）
输出:  渲染契约断言（内容契约/零公式/字节确定性/未校核显式行/流包装异常）
"""

# ══════════════════════════════════════════════════════════════════
# 规格说明（est-20261001 §三 D1~D3——案乙 openpyxl 直写，calcbook 模板
#   机制不套用；渲染件零业务计算，全部数值投递自 EstimateSheet/FeeLine/
#   IndicatorReading 字段）
#
# 【覆盖面】
#   - 内容契约：两工作表（概算汇总/指标校核）；抬头=工况+三元组
#     （design_hash/engine_version/data_version）；分部分项逐行名称列
#     =book.get(price_key).name（渲染名称列单一真源）；合价/小计/
#     grand_total 数值投递；指标行状态 OK|WARN；
#   - 零公式：全部单元格 data_type != "f"（R2——计算单点在 Python）；
#   - 字节确定性：隔秒双渲染字节恒等+core.xml modified==固定纪元定值
#     （calcbook R4/批 14-FIX 同款断言——新工作簿 created 锚定固定纪元）；
#   - 未校核显式行：IndicatorReport(checked=False) → 指标校核表显式
#     「未校核」行（indicators R4 禁静默通过）；
#   - 流包装面（estimate_render_flow——白名单单件承载，flows 镜像件
#     test_flows.py 不在本批白名单）：原子落盘零 .tmp 残留；cost 四域
#     异常包装 InvalidFlowError from exc（消息透传含可用工况集）；'..'
#     路径拒；EstimateFlowResult.book 字段=PriceBook 单一真源。
#   - 回炉轮1（rework-est-20261001-r1）：R5 禁公式注入拒（恶意 '=' 前缀
#     name→域异常——check_no_formulas 共享面）；R7 渲染域异常族
#     （InvalidEstimateRenderError：book 失联键/checked=True 空 readings
#     不可达态——flow 面 InvalidFlowError from exc 链保真）；R8 费用行
#     与指标数值锚扩面（rate/base_amount/amount+value/带上下限）。
# 【替身口径】零替身（golden 计算链+真数据包 _REPO_DATA——test_flows.py
#   estimate 用例同源口径）；dataclasses.replace 仅构造 checked=False 面。
# ══════════════════════════════════════════════════════════════════

from __future__ import annotations

import importlib
import re
import time
import zipfile
from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

_mod = importlib.import_module("waterprint.trace.estimate_sheet")
_flows = importlib.import_module("waterprint.flows")

_REPO_DATA = Path(__file__).resolve().parents[3] / "data"

# 工作表名（渲染契约面——测试与实现双侧共同锚）
_SHEET_SUMMARY = "概算汇总表"
_SHEET_INDICATORS = "指标校核表"
_FIXED_MODIFIED = b"2000-01-01T00:00:00Z"


def _outcome(golden_data_dir: Path):
    """golden estimate 流产物（design 工况——test_flows.py 同源真值链）。"""
    from waterprint.app import load_project

    project = load_project(golden_data_dir / "municipal_34760" / "input_project.json")
    env = _flows.build_env_flow(_REPO_DATA, project)
    conditions = _flows.build_condition_flow(project, None)
    plant = _flows.run_calc_flow(project, conditions, env, ()).plant
    return project, _flows.estimate_summary_flow(
        plant, condition_key="design", data_dir=_REPO_DATA
    )


def _cells_text(path: Path) -> str:
    """全部字符串单元格拼接（锚点断言载体——非全字节断言）。"""
    workbook = load_workbook(path)
    return "\n".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str)
    )


def _numeric_cells(path: Path, sheet_name: str) -> set[float]:
    """单工作表数值单元格集合（数值投递断言载体）。"""
    workbook = load_workbook(path)
    return {
        float(cell.value)
        for row in workbook[sheet_name].iter_rows()
        for cell in row
        if isinstance(cell.value, int | float) and not isinstance(cell.value, bool)
    }


def _assert_number_in(numbers: set[float], value: float) -> None:
    """数值投递断言（xlsx 数值面=%.16g 序列化——16 位有效数字口径，
    calcbook 数值占位符同面先例；approx 容差=末位舍入面）。"""
    assert any(number == pytest.approx(value, rel=1e-15, abs=1e-12)
               for number in numbers), f"数值未投递：{value!r}"


def test_render_estimate_xlsx_content_contract(
    golden_data_dir: Path, tmp_path: Path
) -> None:
    """内容契约：两表+抬头三元组+名称列真源 book+grand_total 投递+零公式。"""
    project, outcome = _outcome(golden_data_dir)
    assert isinstance(outcome.book.get(outcome.sheet.detail_rows[0].price_key).name, str)
    out = tmp_path / "estimate.xlsx"
    assert _mod.render_estimate_xlsx(outcome, out) == out
    workbook = load_workbook(out)
    assert workbook.sheetnames == [_SHEET_SUMMARY, _SHEET_INDICATORS]
    for sheet in workbook.worksheets:  # R2 零公式（计算单点在 Python）
        for row in sheet.iter_rows():
            for cell in row:
                assert cell.data_type != "f"
    text = _cells_text(out)
    repro = outcome.sheet.repro
    assert "design" in text  # 抬头工况
    assert repro.design_hash in text  # 三元组自证（audit 头部同款口径）
    assert repro.engine_version in text
    assert repro.data_version in text
    assert "元" in text  # 金额单位表头注明（批6d 折元口径）
    for row in outcome.sheet.detail_rows:  # 名称列=book.get(price_key).name
        assert outcome.book.get(row.price_key).name in text
    for line in (*outcome.sheet.measure, *outcome.sheet.indirect,
                 *outcome.sheet.reserve, *outcome.sheet.tax):
        assert line.fee_key in text  # 费用行名称（fee_key）
        assert line.base in text  # 基数 DSL 文本
        assert line.source in text  # 出处
    assert "工程概算总计" in text  # grand_total 显式行
    numbers = _numeric_cells(out, _SHEET_SUMMARY)
    for value in (  # 数值投递（非字符串——16 位有效数字口径）
        outcome.sheet.grand_total,
        outcome.sheet.detail_subtotal,
        outcome.sheet.construction_subtotal,
        outcome.sheet.subtotal,
        outcome.sheet.reserve_subtotal,
    ):
        _assert_number_in(numbers, value)
    for line in (*outcome.sheet.measure, *outcome.sheet.indirect,  # R8：费用行三数值
                 *outcome.sheet.reserve, *outcome.sheet.tax):
        _assert_number_in(numbers, line.rate)
        _assert_number_in(numbers, line.base_amount)
        _assert_number_in(numbers, line.amount)
    indicator_numbers = _numeric_cells(out, _SHEET_INDICATORS)
    for reading in outcome.report.readings:  # R8：指标值/带上下限数值锚
        lower, upper = reading.band
        _assert_number_in(indicator_numbers, reading.value)
        _assert_number_in(indicator_numbers, lower)
        _assert_number_in(indicator_numbers, upper)


def test_render_estimate_xlsx_deterministic_across_seconds(
    golden_data_dir: Path, tmp_path: Path
) -> None:
    """字节确定性：隔秒双渲染字节恒等+modified==固定纪元（calcbook R4 同款）。"""
    _, outcome = _outcome(golden_data_dir)
    first = tmp_path / "a.xlsx"
    _mod.render_estimate_xlsx(outcome, first)
    time.sleep(1.1)  # 跨秒——created/modified 时钟面未归一时必现漂移
    second = tmp_path / "b.xlsx"
    _mod.render_estimate_xlsx(outcome, second)
    assert first.read_bytes() == second.read_bytes()
    with zipfile.ZipFile(first) as archive:
        core_xml = archive.read("docProps/core.xml")
    modified = re.search(rb"<dcterms:modified[^>]*>([^<]*)</dcterms:modified>", core_xml)
    assert modified is not None and modified.group(1) == _FIXED_MODIFIED
    created = re.search(rb"<dcterms:created[^>]*>([^<]*)</dcterms:created>", core_xml)
    assert created is not None  # 新工作簿 created 锚定固定纪元（非渲染时钟）
    first_again = tmp_path / "a.xlsx"  # 同名重渲染幂等（覆盖面）
    _mod.render_estimate_xlsx(outcome, first_again)
    assert first_again.read_bytes() == second.read_bytes()


def test_render_estimate_xlsx_unchecked_explicit_row(
    golden_data_dir: Path, tmp_path: Path
) -> None:
    """未校核显式行：checked=False → 指标校核表显式「未校核」（禁静默通过）。"""
    from waterprint.cost.indicators import IndicatorReport

    _, outcome = _outcome(golden_data_dir)
    unchecked = replace(
        outcome, report=IndicatorReport(readings=(), checked=False)
    )
    out = tmp_path / "unchecked.xlsx"
    _mod.render_estimate_xlsx(unchecked, out)
    text = _cells_text(out)
    assert "未校核" in text
    assert _SHEET_INDICATORS in load_workbook(out).sheetnames  # 表恒在（行显式）


def test_estimate_render_flow_atomic_and_wraps_domain_errors(
    golden_data_dir: Path, tmp_path: Path
) -> None:
    """流包装面：原子落盘零 .tmp 残留；cost 域异常包装 InvalidFlowError
    from exc（消息透传含可用工况集）；'..' 路径拒（R2 audit 同口径）。"""
    project, _ = _outcome(golden_data_dir)
    plant = _flows.run_calc_flow(
        project, _flows.build_condition_flow(project, None),
        _flows.build_env_flow(_REPO_DATA, project), (),
    ).plant
    out = tmp_path / "e.xlsx"
    assert _flows.estimate_render_flow(
        project, plant, out, condition_key="design", data_dir=_REPO_DATA
    ) == out
    assert out.is_file() and out.read_bytes()[:2] == b"PK"  # xlsx zip 魔数
    assert not list(tmp_path.glob("*.tmp"))  # 半写 tmp 不留（GR-38+H5）
    from waterprint.cost.takeoff import InvalidTakeoffError

    with pytest.raises(_flows.InvalidFlowError, match="可用工况") as wrapped:
        _flows.estimate_render_flow(
            project, plant, tmp_path / "bad.xlsx",
            condition_key="no_such_condition", data_dir=_REPO_DATA,
        )
    assert isinstance(wrapped.value.__cause__, InvalidTakeoffError)  # from exc 链保真
    assert not list(tmp_path.glob("*.tmp"))  # 异常路径 tmp 同样清理（H5）
    with pytest.raises(_flows.InvalidFlowError, match="\\.\\."):
        _flows.estimate_render_flow(
            project, plant, tmp_path / ".." / "escape.xlsx",
            condition_key="design", data_dir=_REPO_DATA,
        )


def test_estimate_flow_result_carries_price_book(
    golden_data_dir: Path,
) -> None:
    """EstimateFlowResult.book 字段：PriceBook 直投（渲染名称列单一真源，
    services/cost.py R5 同源口径——additive 字段消费面自证）。"""
    from waterprint.cost.prices import PriceBook

    _, outcome = _outcome(golden_data_dir)
    assert isinstance(outcome.book, PriceBook)
    # repro.data_version=UF-10 聚合串（coefficients@…+unit_prices@…）；
    # book.data_version=price_data_version（聚合串成员——estimate.py 注记口径）
    assert outcome.book.data_version in outcome.sheet.repro.data_version


def _book_with_first_name(outcome, malicious_name: str):
    """构造首明细行 name 恶意注入的 book（其余条目原样——R5 注入载体）。"""
    from waterprint.cost.prices import PriceBook

    real = outcome.book
    first = outcome.sheet.detail_rows[0].price_key
    entries = {  # 渲染消费面=分部分项 price_key（名称列真源）——逐行重建
        row.price_key: (
            replace(real.get(row.price_key), name=malicious_name)
            if row.price_key == first else real.get(row.price_key)
        )
        for row in outcome.sheet.detail_rows
    }
    return PriceBook(entries, real.data_version)


def test_render_estimate_xlsx_rejects_formula_injection(
    golden_data_dir: Path, tmp_path: Path
) -> None:
    """R5（回炉轮1）：'=' 前缀 name（openpyxl 按公式存储——真实注入面，
    book name/source/fee_key/base DSL 透传串同守）→ 域异常拒（禁公式
    扫描共享面 check_no_formulas，落盘前调用）。"""
    from waterprint.trace.estimate_sheet import InvalidEstimateRenderError

    _, outcome = _outcome(golden_data_dir)
    bad = replace(outcome, book=_book_with_first_name(outcome, "=SUM(A1:A2)"))
    with pytest.raises(InvalidEstimateRenderError, match="公式"):
        _mod.render_estimate_xlsx(bad, tmp_path / "evil.xlsx")
    assert not (tmp_path / "evil.xlsx").exists()  # 拒于落盘之前


def test_render_estimate_xlsx_rejects_unreachable_empty_readings(
    golden_data_dir: Path, tmp_path: Path
) -> None:
    """R7（回炉轮1 d1-F6）：checked=True 而 readings=() 不可达态显式拒
    （装载面守卫应保证 checked⇒逐带一行——禁静默空校核表通过）。"""
    from waterprint.cost.indicators import IndicatorReport
    from waterprint.trace.estimate_sheet import InvalidEstimateRenderError

    _, outcome = _outcome(golden_data_dir)
    inverted = replace(
        outcome, report=IndicatorReport(readings=(), checked=True)
    )
    with pytest.raises(InvalidEstimateRenderError, match="不可达"):
        _mod.render_estimate_xlsx(inverted, tmp_path / "y.xlsx")


def test_estimate_render_flow_wraps_render_domain_errors(
    golden_data_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """R7（回炉轮1）：渲染域异常（book 失联键）→ flow 面 InvalidFlowError
    from exc 链保真（渲染调用并入包装 try 范围——worker 项级收集面同族）。"""
    from waterprint.cost.prices import PriceBook
    from waterprint.trace.estimate_sheet import InvalidEstimateRenderError

    project, outcome = _outcome(golden_data_dir)
    plant = _flows.run_calc_flow(
        project, _flows.build_condition_flow(project, None),
        _flows.build_env_flow(_REPO_DATA, project), (),
    ).plant
    bad = replace(outcome, book=PriceBook({}, outcome.book.data_version))
    monkeypatch.setattr(
        _flows, "estimate_summary_flow",
        lambda _plant, *, condition_key, data_dir: bad,
    )
    with pytest.raises(_flows.InvalidFlowError, match="失联") as wrapped:
        _flows.estimate_render_flow(
            project, plant, tmp_path / "x.xlsx",
            condition_key="design", data_dir=_REPO_DATA,
        )
    assert isinstance(wrapped.value.__cause__, InvalidEstimateRenderError)
    assert not list(tmp_path.glob("*.tmp"))  # 包装路径 tmp 同样清理（H5）
